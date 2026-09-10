"""
Generador de reportes en Excel (.xlsx) para el Sistema de Viáticos usando openpyxl.
"""
from io import BytesIO
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import utils


def generar_excel_reporte(metricas: dict, viaticos: list, config: dict, filtros: dict, output_target) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reporte de Viáticos"
    ws.views.sheetView[0].showGridLines = True

    # Paleta de colores
    NAVY_HEADER = "1E3A8A"
    LIGHT_GRAY = "F1F5F9"
    ALT_ROW = "F8FAFC"
    BORDER_COLOR = "CBD5E1"
    
    font_univ = Font(name="Calibri", size=13, bold=True, color="1E293B")
    font_fac = Font(name="Calibri", size=11, bold=True, color="334155")
    font_rep_title = Font(name="Calibri", size=12, bold=True, color="1E3A8A")
    font_meta = Font(name="Calibri", size=9, italic=True, color="64748B")
    
    font_card_lbl = Font(name="Calibri", size=8, bold=True, color="475569")
    font_card_val = Font(name="Calibri", size=11, bold=True, color="0F172A")
    
    font_th = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    font_td = Font(name="Calibri", size=9, color="1E293B")
    font_td_bold = Font(name="Calibri", size=9, bold=True, color="1E293B")
    
    fill_header = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
    fill_alt = PatternFill(start_color=ALT_ROW, end_color=ALT_ROW, fill_type="solid")
    fill_card = PatternFill(start_color=LIGHT_GRAY, end_color=LIGHT_GRAY, fill_type="solid")
    fill_total = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    
    thin_border = Border(
        left=Side(style="thin", color=BORDER_COLOR),
        right=Side(style="thin", color=BORDER_COLOR),
        top=Side(style="thin", color=BORDER_COLOR),
        bottom=Side(style="thin", color=BORDER_COLOR)
    )
    
    # 1. Encabezado Institucional
    univ = config.get("universidad", "UNIVERSIDAD NACIONAL DE SAN LUIS")
    fac = config.get("facultad", "FACULTAD DE CIENCIAS ECONÓMICAS, JURÍDICAS Y SOCIALES")
    
    ws.merge_cells("A1:K1")
    ws["A1"] = univ
    ws["A1"].font = font_univ
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    
    ws.merge_cells("A2:K2")
    ws["A2"] = fac
    ws["A2"].font = font_fac
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    
    ws.merge_cells("A3:K3")
    ws["A3"] = "REPORTE OFICIAL DE LIQUIDACIÓN DE VIÁTICOS"
    ws["A3"].font = font_rep_title
    ws["A3"].alignment = Alignment(horizontal="center", vertical="center")
    
    f_desde = utils.formato_fecha(filtros.get("fecha_desde")) or "Inicio"
    f_hasta = utils.formato_fecha(filtros.get("fecha_hasta")) or "Fin"
    estado_f = filtros.get("estado") or "Todos"
    
    ws.merge_cells("A4:K4")
    ws["A4"] = f"Período: {f_desde} al {f_hasta}  |  Estado: {estado_f}  |  Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    ws["A4"].font = font_meta
    ws["A4"].alignment = Alignment(horizontal="center", vertical="center")
    
    # 2. Tarjetas de Resumen Financiero
    # Celdas: Total Solicitudes, Total Monto, Pendientes, Pagados, Rendidos
    cards = [
        ("TOTAL SOLICITUDES", f"{metricas.get('total_cantidad', 0)}", "A6:B6", "A7:B7"),
        ("TOTAL LIQUIDADO", f"{utils.formato_moneda(metricas.get('total_monto', 0))}", "C6:D6", "C7:D7"),
        ("PENDIENTES", f"{metricas.get('pendientes_cantidad', 0)} ({utils.formato_moneda(metricas.get('pendientes_monto', 0))})", "E6:F6", "E7:F7"),
        ("PAGADOS", f"{metricas.get('pagados_cantidad', 0)} ({utils.formato_moneda(metricas.get('pagados_monto', 0))})", "G6:H6", "G7:H7"),
        ("RENDIDOS", f"{metricas.get('rendidos_cantidad', 0)} ({utils.formato_moneda(metricas.get('rendidos_monto', 0))})", "I6:K6", "I7:K7"),
    ]
    
    for title, val, rng_title, rng_val in cards:
        ws.merge_cells(rng_title)
        top_cell = ws[rng_title.split(":")[0]]
        top_cell.value = title
        top_cell.font = font_card_lbl
        top_cell.fill = fill_card
        top_cell.alignment = Alignment(horizontal="center", vertical="center")
        
        ws.merge_cells(rng_val)
        val_cell = ws[rng_val.split(":")[0]]
        val_cell.value = val
        val_cell.font = font_card_val
        val_cell.fill = fill_card
        val_cell.alignment = Alignment(horizontal="center", vertical="center")

    # 3. Tabla de Viáticos
    headers = [
        "N° Viático", "Fecha", "Beneficiario", "Cargo",
        "Salida", "Llegada", "Días", "Valor Diario",
        "Total Liq.", "Imputación", "N° Transferencia", "Expediente", "Estado", "Misión / Destino"
    ]
    
    row_start = 9
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=row_start, column=col_idx, value=h)
        cell.font = font_th
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    
    current_row = row_start + 1
    total_dias = 0.0
    total_importe = 0.0
    
    for v in viaticos:
        is_alt = (current_row % 2 == 0)
        row_fill = fill_alt if is_alt else None
        
        dias = float(v.get("cant_dias") or 0.0)
        v_diario = float(v.get("valor_diario") or 0.0)
        imp = float(v.get("importe_total") or 0.0)
        total_dias += dias
        total_importe += imp
        
        row_values = [
            v.get("nro_viatico", ""),
            utils.formato_fecha(v.get("fecha")),
            v.get("apellido_nombre", ""),
            v.get("cargo_codigo", "") or v.get("cargo_denominacion", ""),
            f"{utils.formato_fecha(v.get('fecha_desde'))} {v.get('hora_desde') or ''}".strip(),
            f"{utils.formato_fecha(v.get('fecha_hasta'))} {v.get('hora_hasta') or ''}".strip(),
            dias,
            v_diario,
            imp,
            v.get("imputacion", ""),
            v.get("cheque", ""),
            v.get("expediente", ""),
            v.get("estado", "Pendiente"),
            f"{v.get('mision', '')} - {v.get('lugar', '')}".strip(" -")
        ]
        
        for col_idx, val in enumerate(row_values, 1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = font_td
            cell.border = thin_border
            if row_fill:
                cell.fill = row_fill
                
            # Alineaciones específicas y formatos
            if col_idx in (1, 2, 5, 6, 11, 12, 13):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx in (7,):
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = "#,##0.0"
            elif col_idx in (8, 9):
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = "$ #,##0.00"
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
                
        current_row += 1

    # Fila de Totales
    total_cells = ws.cell(row=current_row, column=1, value="TOTALES")
    total_cells.font = font_td_bold
    total_cells.alignment = Alignment(horizontal="center", vertical="center")
    total_cells.fill = fill_total
    total_cells.border = thin_border
    
    for c in range(2, 7):
        cell = ws.cell(row=current_row, column=c, value="")
        cell.fill = fill_total
        cell.border = thin_border
        
    cell_dias = ws.cell(row=current_row, column=7, value=total_dias)
    cell_dias.font = font_td_bold
    cell_dias.alignment = Alignment(horizontal="right", vertical="center")
    cell_dias.number_format = "#,##0.0"
    cell_dias.fill = fill_total
    cell_dias.border = thin_border
    
    cell_empty = ws.cell(row=current_row, column=8, value="")
    cell_empty.fill = fill_total
    cell_empty.border = thin_border
    
    cell_tot = ws.cell(row=current_row, column=9, value=total_importe)
    cell_tot.font = font_td_bold
    cell_tot.alignment = Alignment(horizontal="right", vertical="center")
    cell_tot.number_format = "$ #,##0.00"
    cell_tot.fill = fill_total
    cell_tot.border = thin_border
    
    for c in range(10, 15):
        cell = ws.cell(row=current_row, column=c, value="")
        cell.fill = fill_total
        cell.border = thin_border

    # Ajuste automático de anchos de columna
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row < row_start:
                continue
            v_str = str(cell.value or "")
            if len(v_str) > max_len:
                max_len = len(v_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 11)

    wb.save(output_target)
